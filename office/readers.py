"""Bounded file/text extraction and pinned public-URL fetching."""
import http.client
import io
import ipaddress
import json
from pathlib import Path
import re
import socket
import ssl
from html.parser import HTMLParser
from urllib.parse import urlsplit,urlunsplit,urljoin
import zipfile

MAX_BYTES=10_000_000
MAX_TEXT=120000
TEXT_EXT={'.txt','.md','.csv','.tsv','.json','.py','.js','.ts','.tsx','.jsx','.html','.css','.sql','.yml','.yaml','.toml','.ipynb','.log','.xml'}
IMAGE_EXT={'.png','.jpg','.jpeg','.webp'}


def office_zip(data):
    bundle=zipfile.ZipFile(io.BytesIO(data))
    if len(bundle.infolist())>500 or sum(i.file_size for i in bundle.infolist())>30_000_000:
        raise ValueError('압축 내부 자료가 너무 큽니다.')
    return bundle


def decode(data):
    for encoding in ('utf-8-sig','cp949','utf-16'):
        try:return data.decode(encoding)
        except UnicodeError:pass
    raise ValueError('텍스트 인코딩을 읽지 못했습니다. UTF-8로 저장해주세요.')


class VisibleHTML(HTMLParser):
    def __init__(self):super().__init__();self.skip=0;self.parts=[]
    def handle_starttag(self,tag,attrs):
        if tag in ('script','style','noscript'):self.skip+=1
        if tag in ('p','div','li','h1','h2','h3','br','tr'):self.parts.append('\n')
    def handle_endtag(self,tag):
        if tag in ('script','style','noscript') and self.skip:self.skip-=1
    def handle_data(self,data):
        if not self.skip:self.parts.append(data)


def extract(name,data):
    if len(data)>MAX_BYTES:raise ValueError('파일당 최대 10MB입니다.')
    extension=Path(name).suffix.lower()
    note=''
    if extension in IMAGE_EXT:
        from PIL import Image
        image=Image.open(io.BytesIO(data));image.verify()
        return '이미지 자료: '+Path(name).name,'이미지는 Codex의 이미지 입력으로 전달합니다.'
    if extension in TEXT_EXT:
        text=decode(data)
        if extension=='.html':
            parser=VisibleHTML();parser.feed(text);text=''.join(parser.parts)
        if extension=='.ipynb':
            notebook=json.loads(text);text='\n'.join(''.join(cell.get('source',[])) for cell in notebook.get('cells',[]))
    elif extension=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data))
        if reader.is_encrypted:raise ValueError('암호화 PDF는 해제 후 업로드하세요.')
        text='\n'.join(f'[페이지 {number+1}]\n'+(page.extract_text() or '') for number,page in enumerate(reader.pages[:100]))
        if len(reader.pages)>100:note='앞 100페이지만 텍스트를 추출했습니다.'
        if not text.strip() or not re.sub(r'\[페이지 \d+\]','',text).strip():
            raise ValueError('스캔 PDF에서 텍스트를 추출하지 못했습니다. OCR한 PDF 또는 이미지로 올려주세요.')
    elif extension=='.docx':
        office_zip(data).close()
        from docx import Document
        doc=Document(io.BytesIO(data));text='\n'.join(p.text for p in doc.paragraphs)
        text+='\n'+'\n'.join(' | '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
    elif extension=='.xlsx':
        office_zip(data).close()
        import openpyxl
        book=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False)
        try:
            parts=[]
            for sheet in book.worksheets[:10]:
                parts.append('시트: '+sheet.title)
                for row in sheet.iter_rows(max_row=min(sheet.max_row or 200,200),max_col=min(sheet.max_column or 30,30),values_only=True):
                    parts.append(' | '.join('' if v is None else str(v) for v in row))
            text='\n'.join(parts);note='미리보기는 최대 10시트·200행·30열입니다. 전체 원본은 직원 작업에서 읽을 수 있습니다.'
        finally:book.close()
    elif extension=='.zip':
        with office_zip(data) as bundle:
            parts=[]
            for info in bundle.infolist():
                path=Path(info.filename)
                if info.is_dir() or '..' in path.parts or path.is_absolute() or '\\' in info.filename or info.file_size>MAX_BYTES:continue
                if path.suffix.lower() in TEXT_EXT and len(parts)<30:
                    parts.append('파일: '+info.filename+'\n'+decode(bundle.read(info))[:10000])
            text='\n\n'.join(parts);note='압축은 직접 풀지 않습니다. 텍스트 파일 최대 30개를 미리 읽었습니다.'
    else:raise ValueError('지원 형식: 텍스트·코드·PDF·DOCX·XLSX·PNG/JPG/WEBP·ZIP')
    if not text.strip():raise ValueError('읽을 수 있는 본문이 없습니다.')
    if len(text)>MAX_TEXT:note+=' 텍스트 미리보기가 잘렸습니다. 원본 자료는 보관합니다.'
    return text[:MAX_TEXT],note.strip()


def public_address(url):
    parsed=urlsplit(url)
    if parsed.scheme not in ('http','https') or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError('공개 HTTP/HTTPS 주소를 입력하세요.')
    port=parsed.port or (443 if parsed.scheme=='https' else 80)
    if port not in (80,443):raise ValueError('웹 표준 포트의 공개 주소만 읽습니다.')
    host=parsed.hostname.encode('idna').decode()
    addresses=sorted({row[4][0] for row in socket.getaddrinfo(host,port,type=socket.SOCK_STREAM)})
    if not addresses or any(not ipaddress.ip_address(address).is_global for address in addresses):
        raise ValueError('PC 내부·사설망 주소는 링크 수집 대상이 아닙니다. 파일로 올려주세요.')
    return parsed,host,port,addresses[0]


def request(url):
    parsed,host,port,address=public_address(url)
    connection=http.client.HTTPConnection(host,port,timeout=12)
    sock=socket.create_connection((address,port),timeout=12)
    try:
        connection.sock=ssl.create_default_context().wrap_socket(sock,server_hostname=host) if parsed.scheme=='https' else sock
        path=urlunsplit(('','',parsed.path or '/',parsed.query,''))
        connection.request('GET',path,headers={'User-Agent':'MyLittleOffice/0.9','Accept-Encoding':'identity','Accept':'text/html,text/plain,application/pdf,application/json,*/*;q=0.5'})
        response=connection.getresponse()
        body=response.read(MAX_BYTES+1)
        if len(body)>MAX_BYTES:raise ValueError('링크 자료가 10MB를 넘습니다.')
        return response.status,dict((k.lower(),v) for k,v in response.getheaders()),body
    finally:connection.close();sock.close()


def normalize_url(url):
    parsed=urlsplit(url.strip())
    if parsed.hostname=='github.com':
        parts=parsed.path.strip('/').split('/')
        if len(parts)>=5 and parts[2]=='blob':
            return 'https://raw.githubusercontent.com/'+'/'.join(parts[:2]+parts[3:])
        if len(parts)==2:return 'https://raw.githubusercontent.com/'+'/'.join(parts)+'/HEAD/README.md'
    return url.strip()


def read_url(url):
    current=normalize_url(url)
    for _ in range(5):
        status,headers,data=request(current)
        if status in (301,302,303,307,308):
            if not headers.get('location'):raise ValueError('잘못된 리디렉션입니다.')
            current=urljoin(current,headers['location']);continue
        if status in (401,403):raise ValueError('로그인·접근 권한이 필요한 링크입니다. 파일로 올려주세요.')
        if status!=200:raise ValueError('링크 응답 오류: '+str(status))
        if headers.get('content-encoding','').lower() not in ('','identity'):
            raise ValueError('압축 응답을 읽지 못했습니다. 자료를 파일로 올려주세요.')
        content_type=headers.get('content-type','').split(';')[0].lower()
        name=Path(urlsplit(current).path).name or 'link.txt'
        if content_type in ('text/html','application/xhtml+xml'):name='page.html'
        elif content_type=='application/pdf':name='page.pdf'
        elif content_type=='application/json':name='page.json'
        elif content_type in ('image/png','image/jpeg','image/webp'):name='image.'+{'image/png':'png','image/jpeg':'jpg','image/webp':'webp'}[content_type]
        elif not Path(name).suffix:name+='.txt'
        text,note=extract(name,data)
        return {'url':current,'name':name,'data':data,'text':text,'note':note or '공개 페이지 본문 추출. 로그인·JavaScript 실행은 하지 않습니다.'}
    raise ValueError('리디렉션이 너무 많습니다.')
