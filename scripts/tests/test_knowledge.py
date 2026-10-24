import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from office.store import Store
from office import knowledge, readers


class KnowledgeTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.store=Store(self.temp.name)
        self.store.clock_in()

    def test_frozen_profile_and_sources_reach_workers_and_report(self):
        knowledge.save_profile(self.store,'대표 경험: Python 데이터 분석')
        source=knowledge.upload(self.store,'notes.md','고객 분석 기준'.encode())
        root=self.store.submit('lead','고객 데이터 분석을 진행해줘',delegate=True,source_ids=[source])
        frozen=self.store.get(root)['context_json']
        knowledge.save_profile(self.store,'변경된 대표 프로필입니다')
        self.store.claim()
        self.store.finish(root,{'status':'review','output':json.dumps({'tasks':[{'role':'dev','prompt':'고객 데이터를 분석해줘'}]})})
        child=self.store.children(root)[0]
        self.assertEqual(child['context_json'],frozen)
        self.store.claim()
        self.store.finish(child['id'],{'status':'review','output':'분석 완료'})
        self.assertEqual(self.store.children(root)[-1]['context_json'],frozen)
        text,images=knowledge.prepare_inputs(child,Path(self.temp.name)/'run'/'inputs')
        self.assertIn('Python',text)
        self.assertIn('고객 분석 기준',text)
        self.assertEqual(images,[])
        with self.assertRaises(ValueError):self.store.submit('lead','고객 데이터 분석을 진행해줘',root,delegate=True)

    def test_image_originals_and_integrity(self):
        from PIL import Image
        data=io.BytesIO();Image.new('RGB',(2,2)).save(data,format='PNG')
        source=knowledge.upload(self.store,'avatar.png',data.getvalue())
        task=self.store.get(self.store.submit('dev','이미지를 분석해줘',source_ids=[source]))
        text,images=knowledge.prepare_inputs(task,Path(self.temp.name)/'inputs')
        self.assertEqual(len(images),1)
        self.assertEqual(Path(images[0]).read_bytes(),data.getvalue())
        Path(json.loads(task['context_json'])['sources'][0]['filename']).write_bytes(b'changed')
        with self.assertRaises(ValueError):knowledge.prepare_inputs(task,Path(self.temp.name)/'inputs')

    def test_failed_link_visible_and_selected_priority(self):
        with patch('office.knowledge.read_url',side_effect=ValueError('로그인 필요')):
            identifier=knowledge.link(self.store,'https://example.org/private')
        for i in range(21):knowledge.upload(self.store,str(i)+'.txt',str(i).encode(),scope='company')
        context=knowledge.snapshot(self.store,'참고 자료를 확인해줘',selected=[identifier])
        self.assertEqual(len(context['sources']),20)
        self.assertEqual(context['sources'][0]['status'],'failed')
        self.assertIn('로그인 필요',knowledge.render_context(context))

    def test_public_dns_and_redirect_validation(self):
        with patch('socket.getaddrinfo',return_value=[(2,1,6,'',('127.0.0.1',443))]):
            with self.assertRaises(ValueError):readers.public_address('https://example.org')
        with patch('socket.getaddrinfo',return_value=[(2,1,6,'',('93.184.216.34',443))]):
            self.assertEqual(readers.public_address('https://example.org')[3],'93.184.216.34')
        with patch('office.readers.request',side_effect=[(302,{'location':'http://127.0.0.1'},b''),ValueError('private')]) as request:
            with self.assertRaises(ValueError):readers.read_url('https://example.org')
            self.assertEqual(request.call_args.args[0],'http://127.0.0.1')

    def test_html_and_documents(self):
        text,_=readers.extract('page.html',b'<h1>Hello</h1><script>secret()</script><p>world</p>')
        self.assertIn('Hello',text);self.assertNotIn('secret',text)
        from docx import Document
        doc=Document();doc.add_paragraph('owner knowledge');data=io.BytesIO();doc.save(data)
        self.assertIn('owner knowledge',readers.extract('a.docx',data.getvalue())[0])
        from openpyxl import Workbook
        book=Workbook();book.active.append(['Python',42]);data=io.BytesIO();book.save(data)
        self.assertIn('Python | 42',readers.extract('a.xlsx',data.getvalue())[0])
        self.assertEqual(readers.normalize_url('https://github.com/ryemso/my_little_office'),'https://raw.githubusercontent.com/ryemso/my_little_office/HEAD/README.md')

    def test_profile_import_and_project_isolation(self):
        with self.assertRaises(ValueError):knowledge.import_profile(self.store,'[]')
        self.assertEqual(knowledge.inline_links('(https://example.org/notes).'),['https://example.org/notes'])
        knowledge.import_profile(self.store,json.dumps({'format':'my-little-office-profile-v1','profile':'대표의 검증된 초기 경험'}))
        self.assertIn('초기 경험',knowledge.profile(self.store))
        identifier=knowledge.upload(self.store,'private.txt',b'private',project_id='other')
        with self.assertRaises(ValueError):knowledge.snapshot(self.store,'참고 자료 확인',selected=[identifier])

    def test_codex_receives_profile_image_and_raw_paths(self):
        import time
        from types import SimpleNamespace
        from PIL import Image
        from office.engine import run_codex
        knowledge.save_profile(self.store,'대표 핵심 기술은 Pandas입니다')
        data=io.BytesIO();Image.new('RGB',(2,2)).save(data,format='PNG')
        source=knowledge.upload(self.store,'test.png',data.getvalue())
        self.store.submit('dev','이미지의 의미를 분석해줘',source_ids=[source])
        task=self.store.claim()
        captured=[]
        class Input(io.StringIO):
            def close(inner):captured.append(inner.getvalue());super().close()
        def fake(args,**kwargs):
            self.assertIn('--image',args)
            self.assertTrue(Path(args[args.index('--image')+1]).is_file())
            Path(args[args.index('-o')+1]).write_text('모의 이미지 분석 결과')
            kwargs['stdout'].write(json.dumps({'type':'turn.completed'})+'\n');kwargs['stdout'].flush()
            return SimpleNamespace(stdin=Input(),returncode=0,poll=lambda:0)
        with patch('office.browser_modes.prepare',return_value=([],{'catalog_verified':True,'fixture':True})),patch('office.engine.require_login'),patch('office.engine.codex_command',return_value=['codex']),patch('office.engine.subprocess.Popen',side_effect=fake):
            self.assertEqual(run_codex(self.store,task,time.time()+10)['status'],'review')
        self.assertIn('Pandas',captured[0]);self.assertIn('전체 원본 자료 경로',captured[0])
